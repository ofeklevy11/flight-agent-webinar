"""build_excel.py - turns raw Google Flights scans (from gf.js) into the formatted RTL Excel workbook.

Usage:
  python tools/build_excel.py --query data/raw/<name>.query.json --raw data/raw/<name> [--raw data/raw/<name2>] \
         [--out output/<file>.xlsx] [--notes data/raw/<name>.notes.json] [--copy-to "<dir>"]

--raw <prefix>  reads <prefix>.p1.txt, <prefix>.p2.txt ... (the chunks returned by GF.result(i), concatenated),
                or <prefix>.json when it is a single file.
Prints a JSON summary (top 5, stats, problems) for the agent to quote in chat.
All scoring is computed here, from the constants below, so the same input always gives the same score.
"""
import argparse
import glob
import json
import os
import re
import shutil
import sys
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# ----------------------------------------------------------------------------- scoring constants
STOP_PEN = {0: 0, 1: 10, 2: 22, 3: 35}          # per direction
LAY_SHORT_MIN, LAY_SHORT_PEN = 60, 6              # tight connection
LAY_BUCKETS = [(240, 8), (480, 15), (720, 22)]    # (from minutes, points) per layover; highest bucket reached
LAY_NIGHT_MIN, LAY_NIGHT_PEN = 360, 22            # layover of 6h+ that spans 00:00-05:00
DUR_PER_PCT, DUR_CAP = 0.25, 20                   # total travel time above the fastest option
PRICE_PER_PCT, PRICE_CAP = 0.35, 35               # price above the cheapest option
NIGHT_PEN = 3                                     # each take-off / landing between 00:00 and 05:29
AIRPORT_CHANGE_PEN = 10
GROUND_PEN = 10                                   # train / bus leg inside the ticket
LCC_PEN = 4                                       # low-cost carrier: trolley/bag usually extra
LCC = {"W6", "FR", "U2", "VY", "PC", "TO", "HV", "DY", "LS", "F9", "NK", "W4", "W9", "XQ", "BY", "EW", "SK9"}
BANDS = [(80, "טוב", "C6EFCE", "006100"), (65, "כמעט טוב", "FFEB9C", "7F6000"),
         (50, "גבולי", "FBD5B5", "9C4A00"), (0, "לא טוב", "FFC7CE", "9C0006")]
GRAY_FILL, GRAY_FONT = "E7E6E6", "7F7F7F"

HEB_DAY = ["שני", "שלישי", "רביעי", "חמישי", "שישי", "שבת", "ראשון"]
HEB_CITY = {"IST": "איסטנבול", "SAW": "איסטנבול (סבייהה)", "ATH": "אתונה", "VIE": "וינה", "FRA": "פרנקפורט",
            "MUC": "מינכן", "ZRH": "ציריך", "AMS": "אמסטרדם", "CDG": "פריז", "LHR": "לונדון", "WAW": "ורשה",
            "OTP": "בוקרשט", "SOF": "סופיה", "BEG": "בלגרד", "PRG": "פראג", "BUD": "בודפשט", "FCO": "רומא",
            "MXP": "מילאנו", "LCA": "לרנקה", "PFO": "פאפוס", "BCN": "ברצלונה", "MAD": "מדריד", "DXB": "דובאי",
            "AUH": "אבו דאבי", "DOH": "דוחא", "CAI": "קהיר", "AMM": "עמאן", "WAW ": "ורשה", "KRK": "קרקוב",
            "RIX": "ריגה", "HEL": "הלסינקי", "CPH": "קופנהגן", "BRU": "בריסל", "LIS": "ליסבון", "TLL": "טאלין",
            "TBS": "טביליסי", "LJU": "לובליאנה", "ZAG": "זגרב", "TIA": "טירנה", "SKP": "סקופיה", "VNO": "וילנה",
            "GVA": "ז'נבה", "DUS": "דיסלדורף", "HAM": "המבורג", "BER": "ברלין", "JFK": "ניו יורק", "EWR": "ניוארק"}
DISCLAIMER = "המחיר כולל מסים ואגרות לנוסע אחד, ללא מזוודות ותוספות. בחברות לואו קוסט טרולי/מזוודה בדרך כלל בתשלום נוסף."

# ----------------------------------------------------------------------------- helpers
def hm(minutes):
    return f"{minutes // 60}:{minutes % 60:02d}" if minutes is not None else ""


def city(code, cities):
    return HEB_CITY.get(code) or cities.get(code) or code


def fmt_date(d):  # 2026-11-01 -> 01/11/2026 (ראשון)
    dt = datetime.strptime(d, "%Y-%m-%d")
    return f"{dt:%d/%m/%Y} ({HEB_DAY[dt.weekday()]})"


def band(score):
    for lo, word, fill, font in BANDS:
        if score >= lo:
            return word, fill, font


def fnv(text):
    h = 2166136261
    for i in range(0, len(text.encode("utf-16-le")), 2):
        h ^= int.from_bytes(text.encode("utf-16-le")[i:i + 2], "little")
        h = (h * 16777619) & 0xFFFFFFFF
    return h


def load_raw(prefix):
    """Reads <prefix>.p1.txt, .p2.txt ... Each part may be the tool's JSON-string literal copied verbatim
    (with surrounding quotes and backslashes) or the plain text; the joined text must pass the chk test."""
    parts = sorted(glob.glob(prefix + ".p*.txt"), key=lambda p: int(re.search(r"\.p(\d+)\.txt$", p).group(1)))
    if not parts and os.path.exists(prefix + ".json"):
        parts = [prefix + ".json"]
    if not parts:
        raise SystemExit(f"raw not found: {prefix}.p*.txt")
    chunks = []
    for p in parts:
        s = open(p, encoding="utf-8").read().strip()
        if s.startswith('"') and s.endswith('"'):
            s = json.loads(s)
        chunks.append(s)
    text = "".join(chunks)
    m = re.match(r'\{"chk":"([0-9a-f]{8})"', text)
    if m:
        want = int(m.group(1), 16)
        got = fnv(text.replace(m.group(1), "00000000", 1))
        if want != got:
            raise SystemExit(f"הקובץ {prefix} פגום (סכום בדיקה לא תואם). להריץ GF.result(i) שוב ולכתוב את הטקסט כפי שהוא, בלי לשנות.")
    return json.loads(text)


def expand_dir(rec, base_date, names):
    """compact [legs, lay, tot] -> dict with real datetimes (year inferred from base_date)."""
    legs_c, lay_c, tot = rec
    year = int(base_date[:4])
    prev = None
    legs = []
    for fr, to, dm, dt_, am, at_, mins, al, num in legs_c:
        def mk(md, hhmm):
            nonlocal year, prev
            m = int(md[:2])
            if prev is not None and m < prev:
                year += 1
            prev = m
            return datetime(year, m, int(md[2:]), int(hhmm[:2]), int(hhmm[2:]))
        d1 = mk(dm, dt_)
        d2 = mk(am, at_)
        legs.append({"from": fr, "to": to, "dep": d1, "arr": d2, "min": mins, "al": al, "num": num})
    return {"legs": legs, "lay": lay_c, "tot": tot or sum(l["min"] for l in legs),
            "stops": len(legs) - 1, "airlines": list(dict.fromkeys(names.get(l["al"], l["al"]) for l in legs)),
            "nums": " + ".join(l["al"] + " " + l["num"] for l in legs)}


def load_all(raws):
    its, names, cities, places, ats, scans, warns = [], {}, {}, [], [], [], []
    for pre in raws:
        r = load_raw(pre)
        names.update(r.get("names", {}))
        cities.update(r.get("cities", {}))
        if r.get("at"):
            ats.append(r["at"])
        if r.get("place"):
            places.append(r["place"])
        if r["kind"] == "scan":
            y = r.get("year", datetime.now().year)

            def full(md, dep_md=None):
                yy = y + (1 if dep_md and md < dep_md else 0)
                return f"{yy}-{md[:2]}-{md[2:]}"
            r["pairs"] = [[full(d), full(rt, d), mn, ns, n] for d, rt, mn, ns, n in r["pairs"]]
            r["failed"] = [[full(d), full(rt, d)] for d, rt in r.get("failed", [])]
            scans.append(r)
            warns += [f"סריקה: {len(r['failed'])} צמדי תאריכים נכשלו" ] if r.get("failed") else []
        for e in r.get("exacts", []):
            outs = [expand_dir(v, e["dep"], names) for v in e["out"]]
            rets = [expand_dir(v, e["ret"], names) for v in e["ret_"]]
            warns += [f"{e['dep']}..{e['ret']}: {w}" for w in e.get("warn", [])]
            seen = {}
            for oid, rid, price in e["itins"]:  # same out+ret can repeat (fare variants): keep the cheapest
                k = (oid, rid)
                if k in seen and (price is None or (seen[k] is not None and seen[k] <= price)):
                    continue
                seen[k] = price
            for (oid, rid), price in seen.items():
                its.append({"dep": e["dep"], "ret": e["ret"], "oid": (e["dep"], oid), "rid": (e["ret"], rid) if rid >= 0 else None,
                            "out": outs[oid], "back": rets[rid] if rid >= 0 else None, "price": price, "pairkey": (e["dep"], e["ret"])})
    return its, names, cities, scans, warns, ats


# ----------------------------------------------------------------------------- scoring
def lay_flags(d, cities):
    """per-layover list of (minutes, airport, label, penalty, reason)"""
    out = []
    for i, (mins, apt_in, apt_out) in enumerate(d["lay"]):
        pen, why = 0, ""
        a_arr, b_dep = d["legs"][i]["arr"], d["legs"][i + 1]["dep"]
        night = False
        t = a_arr
        while t < b_dep:  # does the wait include 00:00-05:00?
            if t.hour < 5:
                night = True
                break
            t += timedelta(minutes=30)
        if night and mins >= LAY_NIGHT_MIN:
            pen, why = LAY_NIGHT_PEN, "המתנת לילה"
        else:
            for lo, p in LAY_BUCKETS:
                if mins >= lo:
                    pen, why = p, "המתנה ארוכה"
            if mins < LAY_SHORT_MIN:
                pen, why = LAY_SHORT_PEN, "קונקשן הדוק"
        out.append((mins, apt_in, why, pen, apt_in != apt_out))
    return out


def is_night(dt):
    return dt.hour < 5 or (dt.hour == 5 and dt.minute < 30)


def score_all(its, names, cities):
    priced = [x for x in its if x["price"] is not None and x["back"] is not None]
    if not priced:
        for x in its:
            x.update(score=None, ded=[], summary="אין מחיר זמין", lays=[])
        return
    cheapest = min(x["price"] for x in priced)
    fastest = min(x["out"]["tot"] + x["back"]["tot"] for x in priced)
    for x in its:
        x["lays"] = {"out": lay_flags(x["out"], cities), "ret": lay_flags(x["back"], cities) if x["back"] else []}
        if x["price"] is None or x["back"] is None:
            x.update(score=None, ded=[], summary="אין מחיר זמין")
            continue
        ded = []  # (category, points, text)
        for side, label in (("out", "הלוך"), ("back", "חזור")):
            d = x[side]
            n = min(d["stops"], 3)
            if n:
                where = ", ".join(city(l["to"], cities) for l in d["legs"][:-1])
                ded.append(("עצירות", STOP_PEN[n], f"{label}: {'עצירה' if n == 1 else str(n) + ' עצירות'} ({where})"))
            for (mins, apt, why, pen, change) in x["lays"]["out" if side == "out" else "ret"]:
                if pen:
                    ded.append(("המתנות", pen, f"{label}: {why} של {hm(mins)} ש' ב{city(apt, cities)}"))
                if change:
                    ded.append(("אחר", AIRPORT_CHANGE_PEN, f"{label}: החלפת שדה תעופה בעצירה ב{city(apt, cities)}"))
            for l in d["legs"]:
                if re.search(r"rail|train|bus", names.get(l["al"], ""), re.I):
                    ded.append(("אחר", GROUND_PEN, f"{label}: קטע קרקעי (רכבת/אוטובוס)"))
                    break
        total_min = x["out"]["tot"] + x["back"]["tot"]
        ex = (total_min / fastest - 1) * 100
        if ex >= 4:
            ded.append(("משך", round(min(DUR_CAP, DUR_PER_PCT * ex)), f"משך כולל {hm(total_min)} ש', ארוך ב-{round(ex)}% מהמהיר ביותר"))
        px = (x["price"] / cheapest - 1) * 100
        if px >= 3:
            ded.append(("מחיר", round(min(PRICE_CAP, PRICE_PER_PCT * px)), f"יקרה ב-{round(px)}% מהזולה ביותר"))
        nights = [lbl for lbl, dt in (("המראה הלוך", x["out"]["legs"][0]["dep"]), ("נחיתה הלוך", x["out"]["legs"][-1]["arr"]),
                                      ("המראה חזור", x["back"]["legs"][0]["dep"]), ("נחיתה חזור", x["back"]["legs"][-1]["arr"]))
                  if is_night(dt)]
        if nights:
            ded.append(("שעות לילה", NIGHT_PEN * len(nights), "שעות לא נוחות: " + ", ".join(nights)))
        if any(l["al"] in LCC for d in (x["out"], x["back"]) for l in d["legs"]):
            ded.append(("אחר", LCC_PEN, "לואו קוסט: מזוודה/טרולי בתוספת"))
        ded = [d for d in ded if d[1] > 0]
        x["ded"] = ded
        x["score"] = max(0, 100 - sum(p for _, p, _ in ded))
        x["summary"] = " · ".join(f"{t} (−{p})" for _, p, t in sorted(ded, key=lambda z: -z[1])) or "ישירה, זולה ונוחה, בלי חיסרון בולט"
    return cheapest, fastest


def pick_top(scored, n):
    """Best n by score, but varied: one per date pair (several pairs) or one per outbound flight (a single pair),
    then filled with the next best, so the list is not five copies of the same trip."""
    multi = len({x["pairkey"] for x in scored}) > 1
    key = (lambda x: x["pairkey"]) if multi else (lambda x: x["oid"])
    chosen, seen = [], set()
    for x in scored:
        if key(x) not in seen:
            seen.add(key(x))
            chosen.append(x)
        if len(chosen) == n:
            break
    for x in scored:
        if len(chosen) >= n:
            break
        if x not in chosen:
            chosen.append(x)
    return sorted(chosen, key=lambda x: (-x["score"], x["price"]))


# ----------------------------------------------------------------------------- excel plumbing
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F3864")


def mkfill(c):
    return PatternFill("solid", fgColor=c)


def sheet(wb, title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return ws


def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, row, col, val, fill=None, font_color=None, bold=False, fmt=None, wrap=False, align="center", link=None):
    c = ws.cell(row=row, column=col, value=val)
    c.font = Font(name="Arial", size=10, bold=bold, color=font_color or "000000", underline="single" if link else None)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = BORDER
    if fill:
        c.fill = mkfill(fill)
    if fmt:
        c.number_format = fmt
    if link:
        c.hyperlink = link
    return c


def row_style(x):
    if x["score"] is None:
        return GRAY_FILL, GRAY_FONT, "אין מחיר"
    w, f, fc = band(x["score"])
    return f, fc, w


def dir_cells(d, cities):
    first, last = d["legs"][0], d["legs"][-1]
    plus = (last["arr"].date() - first["dep"].date()).days
    arr = last["arr"].strftime("%H:%M") + (f" (+{plus})" if plus > 0 else "")
    if d["stops"] == 0:
        stops = "ישירה"
    else:
        stops = f"{d['stops']} עצירות: " + ", ".join(city(l["to"], cities) for l in d["legs"][:-1])
    lay = ", ".join(f"{hm(m)} ב{city(a, cities)}" for m, a, _ in d["lay"]) or "-"
    return [first["dep"].strftime("%H:%M"), arr, ", ".join(d["airlines"]), d["nums"], stops, hm(d["tot"]), lay]


def gf_link(q, dep, ret):
    from urllib.parse import quote
    t = f"Flights from {q.get('origin_name', 'Tel Aviv')} to {q.get('dest_name', q.get('dest', ''))} on {dep} through {ret}"
    return f"https://www.google.com/travel/flights?q={quote(t)}&curr=ILS&hl=en"


# ----------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--query", required=True)
    ap.add_argument("--raw", action="append", required=True)
    ap.add_argument("--out")
    ap.add_argument("--notes")
    ap.add_argument("--copy-to")
    a = ap.parse_args()
    q = json.load(open(a.query, encoding="utf-8"))
    its, names, cities, scans, warns, ats = load_all(a.raw)
    stats = score_all(its, names, cities)
    problems = list(q.get("problems", [])) + warns
    notes = json.load(open(a.notes, encoding="utf-8")) if a.notes and os.path.exists(a.notes) else {}
    when = max(ats) if ats else datetime.now().isoformat()
    try:
        stamp = datetime.fromisoformat(when.replace("Z", "+00:00")).astimezone().strftime("%d/%m/%Y %H:%M")
    except Exception:
        stamp = when
    dest_label = q.get("dest_label") or q.get("dest_name") or q.get("dest")
    title = f"{q.get('origin_label', 'תל אביב')} ↔ {dest_label}"

    scored = sorted([x for x in its if x["score"] is not None], key=lambda x: (-x["score"], x["price"]))
    unpriced = [x for x in its if x["score"] is None]
    rows = scored + unpriced
    top5 = pick_top(scored, 5)
    cheapest_it = min(scored, key=lambda x: x["price"]) if scored else None
    fastest_it = min(scored, key=lambda x: x["out"]["tot"] + x["back"]["tot"]) if scored else None

    pair_rows = {}
    for s in scans:
        for dep, ret, mn, mns, n in s["pairs"]:
            pair_rows[(dep, ret)] = {"min": mn, "ns": mns, "n": n, "deep": False, "ver": ""}
    exact_min = {}
    for x in its:
        if x["price"] is not None:
            exact_min[x["pairkey"]] = min(x["price"], exact_min.get(x["pairkey"], x["price"]))
    for key, mn in exact_min.items():
        pr = pair_rows.setdefault(key, {"min": mn, "ns": None, "n": None, "deep": True, "ver": ""})
        pr["deep"] = True
        if pr["min"] is None:
            pr["min"] = mn
        elif pr["min"] == mn:
            pr["ver"] = f"אומת: {mn:,} ₪"
        else:
            pr["ver"] = f"שונה: בסריקה {pr['min']:,} ₪, בעמוד התוצאות {mn:,} ₪"
            problems.append(f"אימות {key[0]}..{key[1]}: בסריקה {pr['min']:,} ₪ אבל בעמוד התוצאות {mn:,} ₪ (נעשה שימוש במחיר העמוד)")
            pr["min"] = mn

    wb = Workbook()
    # ---------------------------------------------------------------- summary
    ws = sheet(wb, "סיכום", first=True)
    for col, w in zip("ABCDEFGH", (22, 12, 12, 12, 38, 30, 30, 46)):
        ws.column_dimensions[col].width = w
    ws["A1"] = f"סוכן טיסות: {title}"
    ws["A1"].font = Font(name="Arial", size=18, bold=True, color="1F3864")
    ws.row_dimensions[1].height = 30
    info = [("סוג החיפוש", q.get("skill_label", "")), ("בקשת המשתמש", q.get("user_text", "")),
            ("יציאה מ", q.get("origin_label", "תל אביב")), ("יעד", dest_label),
            ("תאריכים", q.get("dates_label", "")), ("נוסעים / מחלקה", "נוסע אחד, מחלקת תיירים"),
            ("מטבע", "שקלים (₪)"), ("מחירים נכון ל", stamp),
            ("מקור הנתונים", "Google Flights דרך הדפדפן, בלי API ובלי מפתחות"), ("הערה על המחיר", DISCLAIMER)]
    r = 3
    for k, v in info:
        put(ws, r, 1, k, fill="D9E2F3", bold=True, align="right")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        put(ws, r, 2, v, align="right", wrap=True)
        r += 1
    r += 1
    put(ws, r, 1, "מספרים עיקריים", bold=True, fill="D9E2F3", align="right")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    r += 1
    facts = [("צירופי הלוך+חזור שנאספו", len(its)), ("מהם עם מחיר וציון", len(scored)), ("ללא מחיר זמין (אפור)", len(unpriced))]
    if scans:
        facts.append(("צמדי תאריכים שנסרקו", sum(len(s["pairs"]) for s in scans)))
    if cheapest_it:
        facts += [("הכי זול", f"{cheapest_it['price']:,} ₪ ({fmt_date(cheapest_it['dep'])} ← {fmt_date(cheapest_it['ret'])}), ציון {cheapest_it['score']}"),
                  ("הכי מהיר", f"{hm(fastest_it['out']['tot'] + fastest_it['back']['tot'])} ש' בסך הכול, {fastest_it['price']:,} ₪, ציון {fastest_it['score']}"),
                  ("תמורה הכי טובה", f"{top5[0]['price']:,} ₪, ציון {top5[0]['score']} ({top5[0]['out']['nums']} / {top5[0]['back']['nums']})")]
    for k, v in facts:
        put(ws, r, 1, k, align="right", bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        put(ws, r, 2, v, align="right")
        r += 1
    if notes.get("recommendation"):
        r += 1
        put(ws, r, 1, "המלצת קלוד", bold=True, fill="FFF2CC", align="right")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        put(ws, r, 2, notes["recommendation"], fill="FFF2CC", align="right", wrap=True)
        ws.row_dimensions[r].height = 60
        r += 1
    r += 1
    put(ws, r, 1, "5 האפשרויות הטובות ביותר (לפי ציון תמורה למחיר)", bold=True, fill="D9E2F3", align="right")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    r += 1
    header(ws, r, ["#", "ציון", "דירוג", "מחיר ₪", "תאריכים", "הלוך", "חזור", "סיכום קלוד"])
    r += 1
    t5notes = notes.get("top5", [])
    for i, x in enumerate(top5, 1):
        fill, fc, word = row_style(x)
        vals = [i, x["score"], word, x["price"], f"{fmt_date(x['dep'])} ← {fmt_date(x['ret'])}",
                f"{x['out']['legs'][0]['dep']:%H:%M}→{x['out']['legs'][-1]['arr']:%H:%M} {x['out']['nums']}",
                f"{x['back']['legs'][0]['dep']:%H:%M}→{x['back']['legs'][-1]['arr']:%H:%M} {x['back']['nums']}",
                (t5notes[i - 1] + " | " if i - 1 < len(t5notes) else "") + x["summary"]]
        for c, v in enumerate(vals, 1):
            put(ws, r, c, v, fill=fill, font_color=fc, bold=c in (2, 3), fmt='#,##0 "₪"' if c == 4 else None, wrap=c >= 5, align="right" if c >= 5 else "center")
        ws.row_dimensions[r].height = 48
        r += 1
    r += 1
    put(ws, r, 1, "מקרא: איך נקבע הציון", bold=True, fill="D9E2F3", align="right")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    r += 1
    put(ws, r, 1, "הציון מתחיל ב-100 ויורדות נקודות. צבע השורה + מילה + מספר, כדי שיהיה ברור גם בלי צבעים. אפור = אין מחיר זמין, בלי ציון.", align="right", wrap=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 30
    r += 1
    for lo, word, fill, fc in BANDS:
        hi = {80: "80 עד 100", 65: "65 עד 79", 50: "50 עד 64", 0: "0 עד 49"}[lo]
        put(ws, r, 1, word, fill=fill, font_color=fc, bold=True)
        put(ws, r, 2, hi, fill=fill, font_color=fc)
        r += 1
    put(ws, r, 1, "אין מחיר", fill=GRAY_FILL, font_color=GRAY_FONT, bold=True)
    put(ws, r, 2, "בלי ציון", fill=GRAY_FILL, font_color=GRAY_FONT)
    r += 2
    header(ws, r, ["רכיב", "מה מוריד", "נקודות", "", "", "", "", ""])
    r += 1
    legend = [("עצירות", "בכל כיוון בנפרד", f"עצירה אחת {STOP_PEN[1]} · שתיים {STOP_PEN[2]} · שלוש ויותר {STOP_PEN[3]}"),
              ("המתנה בין טיסות", "לכל קונקשן", f"הדוק (מתחת ל-{LAY_SHORT_MIN} דק') {LAY_SHORT_PEN} · 4 עד 8 ש' {LAY_BUCKETS[0][1]} · 8 עד 12 ש' {LAY_BUCKETS[1][1]} · 12+ ש' {LAY_BUCKETS[2][1]} · המתנת לילה (6+ ש' בין 00:00 ל-05:00) {LAY_NIGHT_PEN}"),
              ("משך כולל", "מול הצירוף המהיר ביותר", f"{DUR_PER_PCT} לכל אחוז איטי יותר, עד {DUR_CAP}"),
              ("מחיר", "מול הצירוף הזול ביותר", f"{PRICE_PER_PCT} לכל אחוז יקר יותר, עד {PRICE_CAP}"),
              ("שעות לילה", "המראה או נחיתה בין 00:00 ל-05:29", f"{NIGHT_PEN} לכל אירוע"),
              ("החלפת שדה", "שדה שונה בנחיתה ובהמראה הבאה", f"{AIRPORT_CHANGE_PEN}"),
              ("קטע קרקעי", "רכבת/אוטובוס בתוך הכרטיס", f"{GROUND_PEN}"),
              ("לואו קוסט", "מזוודה/טרולי בדרך כלל בתוספת", f"{LCC_PEN} (פעם אחת לצירוף)")]
    for k, v, p in legend:
        put(ws, r, 1, k, bold=True, align="right")
        put(ws, r, 2, v, align="right", wrap=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        put(ws, r, 3, p, align="right", wrap=True)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    put(ws, r, 1, "בעיות והערות", bold=True, fill="FCE4D6", align="right")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    r += 1
    for p in problems or ["לא נתקלנו בבעיות בריצה הזאת"]:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        put(ws, r, 1, p, align="right", wrap=True)
        r += 1

    # ---------------------------------------------------------------- all flights
    ws = sheet(wb, "כל הטיסות")
    cols = ["#", "ציון", "דירוג", "מחיר כולל ₪", "יציאה", "חזרה"] + [f"הלוך: {t}" for t in ("המראה", "נחיתה", "חברה", "טיסות", "עצירות", "משך", "המתנות")] + \
           [f"חזור: {t}" for t in ("המראה", "נחיתה", "חברה", "טיסות", "עצירות", "משך", "המתנות")] + ["משך כולל", "סיכום קלוד: למה ירדו נקודות", "קישור"]
    widths = [5, 8, 11, 13, 17, 17] + [9, 11, 16, 18, 24, 8, 22] * 2 + [10, 60, 9]
    header(ws, 1, cols, widths)
    for i, x in enumerate(rows, 1):
        fill, fc, word = row_style(x)
        r = i + 1
        back = dir_cells(x["back"], cities) if x["back"] else [""] * 7
        tot = hm(x["out"]["tot"] + x["back"]["tot"]) if x["back"] else ""
        vals = [i, x["score"] if x["score"] is not None else "", word, x["price"] if x["price"] is not None else "אין מחיר",
                fmt_date(x["dep"]), fmt_date(x["ret"])] + dir_cells(x["out"], cities) + back + [tot, x["summary"], "פתח"]
        for c, v in enumerate(vals, 1):
            put(ws, r, c, v, fill=fill, font_color=fc, bold=c in (2, 3), fmt='#,##0 "₪"' if c == 4 and isinstance(v, int) else None,
                wrap=c in (11, 18, 20, 21, 22) or c == len(vals) - 1, align="right" if c in (11, 18, 21, 22) else "center",
                link=gf_link(q, x["dep"], x["ret"]) if c == len(vals) else None)
        ws.row_dimensions[r].height = 44
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    # ---------------------------------------------------------------- one-direction sheets
    def direction_sheet(name, side):
        ws = sheet(wb, name)
        header(ws, 1, ["תאריך", "המראה", "נחיתה", "חברה", "טיסות", "עצירות", "משך", "המתנות", "צירופים", "המחיר הזול ביותר כולל ₪", "הציון הטוב ביותר", "דירוג"],
               [18, 9, 11, 18, 22, 26, 8, 24, 10, 16, 12, 11])
        groups = {}
        for x in its:
            d = x[side]
            if d is None:
                continue
            key = (x["dep"] if side == "out" else x["ret"], x["oid"] if side == "out" else x["rid"])
            g = groups.setdefault(key, {"d": d, "n": 0, "price": None, "score": None})
            g["n"] += 1
            if x["price"] is not None:
                g["price"] = x["price"] if g["price"] is None else min(g["price"], x["price"])
            if x["score"] is not None:
                g["score"] = x["score"] if g["score"] is None else max(g["score"], x["score"])
        items = sorted(groups.items(), key=lambda kv: (kv[1]["score"] is None, -(kv[1]["score"] or 0), kv[1]["price"] or 1e9))
        for r, ((date, _), g) in enumerate(items, 2):
            if g["score"] is None:
                fill, fc, word = GRAY_FILL, GRAY_FONT, "אין מחיר"
            else:
                word, fill, fc = band(g["score"])
            dc = dir_cells(g["d"], cities)
            vals = [fmt_date(date), dc[0], dc[1], dc[2], dc[3], dc[4], dc[5], dc[6], g["n"],
                    g["price"] if g["price"] is not None else "אין מחיר", g["score"] if g["score"] is not None else "", word]
            for c, v in enumerate(vals, 1):
                put(ws, r, c, v, fill=fill, font_color=fc, bold=c in (11, 12), fmt='#,##0 "₪"' if c == 10 and isinstance(v, int) else None,
                    wrap=c in (6, 8), align="right" if c in (4, 5, 6, 8) else "center")
            ws.row_dimensions[r].height = 30
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:L{len(items) + 1}"

    direction_sheet("כיוון הלוך", "out")
    direction_sheet("כיוון חזור", "back")

    # ---------------------------------------------------------------- date scan
    ws = sheet(wb, "סריקת תאריכים")
    header(ws, 1, ["#", "יציאה", "חזרה", "לילות", "המחיר הזול ביותר ₪", "הזול עם הלוך ישיר ₪", "פער מהזול ביותר", "העמקה ואימות", "קישור"], [5, 19, 19, 8, 18, 18, 14, 34, 9])
    plist = sorted(pair_rows.items(), key=lambda kv: (kv[1]["min"] is None, kv[1]["min"] or 0))
    lo = next((v["min"] for _, v in plist if v["min"] is not None), None)
    for i, ((dep, ret), v) in enumerate(plist, 1):
        r = i + 1
        nights = (datetime.strptime(ret, "%Y-%m-%d") - datetime.strptime(dep, "%Y-%m-%d")).days
        if v["min"] is None or lo is None:
            fill, fc, gap = GRAY_FILL, GRAY_FONT, "אין מחיר"
        else:
            pct = (v["min"] / lo - 1) * 100
            gap = f"+{round(pct)}%" if pct >= 1 else "הזול ביותר"
            _, fill, fc = band(100 if pct <= 5 else 70 if pct <= 15 else 55 if pct <= 30 else 0)
        vals = [i, fmt_date(dep), fmt_date(ret), nights, v["min"] if v["min"] is not None else "אין מחיר", v["ns"] if v["ns"] else "-", gap, (v["ver"] or "נבדק לעומק") if v["deep"] else "", "פתח"]
        for c, val in enumerate(vals, 1):
            put(ws, r, c, val, fill=fill, font_color=fc, fmt='#,##0 "₪"' if c in (5, 6) and isinstance(val, int) else None, bold=c == 5,
                link=gf_link(q, dep, ret) if c == 9 else None)
    put(ws, len(plist) + 3, 1, "הצבע לפי הפער מהמחיר הזול ביותר בסריקה: ירוק עד 5%, צהוב עד 15%, כתום עד 30%, אדום מעבר.", align="right")
    ws.merge_cells(start_row=len(plist) + 3, start_column=1, end_row=len(plist) + 3, end_column=9)
    ws.freeze_panes = "A2"
    if len(plist) > 1:  # heat map: departures x returns
        deps = sorted({d for d, _ in pair_rows})
        rets = sorted({r_ for _, r_ in pair_rows})
        hm_ws = sheet(wb, "מפת תאריכים")
        hm_ws.column_dimensions["A"].width = 18
        put(hm_ws, 1, 1, "יציאה ↓ / חזרה ←", fill="1F3864", font_color="FFFFFF", bold=True)
        for j, rt in enumerate(rets, 2):
            put(hm_ws, 1, j, fmt_date(rt)[:5] + "\n" + fmt_date(rt)[-6:-1], fill="1F3864", font_color="FFFFFF", bold=True, wrap=True)
            hm_ws.column_dimensions[get_column_letter(j)].width = 9
        hm_ws.row_dimensions[1].height = 34
        for i2, dp in enumerate(deps, 2):
            put(hm_ws, i2, 1, fmt_date(dp), fill="D9E2F3", bold=True)
            for j, rt in enumerate(rets, 2):
                v = pair_rows.get((dp, rt))
                if not v or v["min"] is None:
                    put(hm_ws, i2, j, "", fill="F2F2F2")
                    continue
                pct = (v["min"] / lo - 1) * 100
                _, fill, fc = band(100 if pct <= 5 else 70 if pct <= 15 else 55 if pct <= 30 else 0)
                put(hm_ws, i2, j, v["min"], fill=fill, font_color=fc, fmt="#,##0")
        hm_ws.freeze_panes = "B2"

    # ---------------------------------------------------------------- value for money
    ws = sheet(wb, "תמורה למחיר")
    header(ws, 1, ["#", "ציון", "דירוג", "מחיר ₪", "יציאה", "חזרה", "הלוך", "חזור", "עצירות (−)", "המתנות (−)", "משך (−)", "מחיר (−)", "שעות לילה (−)", "אחר (−)", "סיכום קלוד"],
           [5, 8, 11, 12, 17, 17, 24, 24, 11, 11, 10, 10, 12, 10, 60])
    for i, x in enumerate(scored[:40], 1):
        fill, fc, word = row_style(x)
        by = {}
        for cat, p, _ in x["ded"]:
            by[cat] = by.get(cat, 0) + p
        vals = [i, x["score"], word, x["price"], fmt_date(x["dep"]), fmt_date(x["ret"]),
                f"{x['out']['legs'][0]['dep']:%H:%M}→{x['out']['legs'][-1]['arr']:%H:%M} {x['out']['nums']}",
                f"{x['back']['legs'][0]['dep']:%H:%M}→{x['back']['legs'][-1]['arr']:%H:%M} {x['back']['nums']}",
                by.get("עצירות", 0), by.get("המתנות", 0), by.get("משך", 0), by.get("מחיר", 0), by.get("שעות לילה", 0), by.get("אחר", 0), x["summary"]]
        for c, v in enumerate(vals, 1):
            put(ws, i + 1, c, v, fill=fill, font_color=fc, bold=c in (2, 3), fmt='#,##0 "₪"' if c == 4 else None, wrap=c == 15, align="right" if c in (7, 8, 15) else "center")
        ws.row_dimensions[i + 1].height = 34
    ws.freeze_panes = "A2"

    # ---------------------------------------------------------------- save
    out = a.out or os.path.join("output", f"{q.get('file_stem', 'flights')}.xlsx")
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    if a.copy_to:
        os.makedirs(a.copy_to, exist_ok=True)
        shutil.copy2(out, a.copy_to)

    def short(x):
        return {"rank_score": x["score"], "word": band(x["score"])[0], "price_ils": x["price"], "dep": x["dep"], "ret": x["ret"],
                "out": f"{x['out']['legs'][0]['dep']:%H:%M}→{x['out']['legs'][-1]['arr']:%H:%M} {x['out']['nums']} ({', '.join(x['out']['airlines'])}, {'ישירה' if x['out']['stops'] == 0 else str(x['out']['stops']) + ' עצירות'}, {hm(x['out']['tot'])})",
                "back": f"{x['back']['legs'][0]['dep']:%H:%M}→{x['back']['legs'][-1]['arr']:%H:%M} {x['back']['nums']} ({', '.join(x['back']['airlines'])}, {'ישירה' if x['back']['stops'] == 0 else str(x['back']['stops']) + ' עצירות'}, {hm(x['back']['tot'])})",
                "why": x["summary"]}
    print(json.dumps({"file": os.path.abspath(out), "stamp": stamp, "itineraries": len(its), "scored": len(scored), "unpriced": len(unpriced),
                      "pairs_scanned": sum(len(s["pairs"]) for s in scans), "cheapest": short(cheapest_it) if cheapest_it else None,
                      "top5": [short(x) for x in top5], "problems": problems}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
